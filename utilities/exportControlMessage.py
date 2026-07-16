from flowstudio import flow_conf_list
from openpyxl import *

# print(flow_conf_list.AUDIO_OBJECT)
wb = load_workbook('msg.xlsx')
sheets = wb.sheetnames
for sheet in sheets:
    worksheet = wb[sheet]
    wb.remove(worksheet)

wb.create_sheet('AO params')
actSheet = wb.active
actSheet['A1'].value = 'Audio Object'
actSheet['B1'].value = 'Parameter Name'
actSheet['C1'].value = 'Min'
actSheet['D1'].value = 'Max'
actSheet['E1'].value = 'Note'

row = 2
col = 2
AO_start_row = 0
param_count = 0

for key, value in flow_conf_list.AUDIO_OBJECT.items():
    print('AO:', key)
    param_count = len(value.items())
    AO_start_row = row

    for k, v in value.items():
        print('[', k, ']')

        cellAO = 'A' + str(row)
        actSheet[cellAO].value = key

        cellParam = 'B' + str(col)
        actSheet[cellParam].value = k
        if k == 'tap':

            # print("------------- tap ----------------")
            valTap = v['cParameter']
            param_count += len(valTap)
            # print('-----------', valTap, '-------------------')
            for j in valTap:
                # print('***************', j, '***************')
                cellAO = 'A' + str(row)
                actSheet[cellAO].value = key
                cellParam = 'B' + str(col)
                actSheet[cellParam].value = j
                jv = valTap[j]
                try:
                    cellMin = 'C' + str(col)
                    val = jv["pMin"]
                    if val == 'off':
                        actSheet[cellMin].value = 0
                    else:
                        actSheet[cellMin].value = val
                except:
                    print('no min value')
                try:
                    cellMax = 'D' + str(col)
                    val = jv["pMax"]
                    if val == 'on':
                        actSheet[cellMax].value = 1
                    else:
                        actSheet[cellMax].value = val
                except:
                    print('no max value')
                val = ''
                try:
                    cellNote = 'E' + str(col)

                    for item in jv["pList"]:
                        val += item + ", "
                    actSheet[cellNote].value = val
                    # print(v['pList'])
                except:
                    print('no List value')

                row += 1
                col += 1

        try:
            cellMin = 'C' + str(col)
            val = v["pMin"]
            if val == 'off':
                actSheet[cellMin].value = 0
            else:
                actSheet[cellMin].value = val
        except:
            print('no min value')

        try:
            cellMax = 'D' + str(col)
            val = v["pMax"]
            if val == 'on':
                actSheet[cellMax].value = 1
            else:
                actSheet[cellMax].value = val

        except:
            print('no max value')

        val = ''
        try:
            cellNote = 'E' + str(col)

            for item in v["pList"]:
                val += item + ", "
            actSheet[cellNote].value = val
            # print(v['pList'])
        except:
            print('no List value')

        row += 1
        col += 1

    # merge the AO cell
    merge_start = 'A' + str(AO_start_row)
    merge_end = 'A' + str(AO_start_row + param_count)

    # try:
    #     actSheet.merged_cells(merge_start + ':' + merge_end)
    #     # print('param count:', param_count)
    #
    # except:
    #     print('no need to merge')

# adjust col width
for col in actSheet.columns:
     max_length = 0
     column = col[0].column_letter # Get the column name
     for cell in col:
         try: # Necessary to avoid error on empty cells
             if len(str(cell.value)) > max_length:
                 max_length = len(str(cell.value))
         except:
             pass
     adjusted_width = (max_length + 2) * 1.1
     actSheet.column_dimensions[column].width = adjusted_width

wb.save('msg.xlsx')
